"""
Generator Surat Penunjukan BNPB
- Dilindungi password
- Dasar Hukum dinamis (bisa tambah / kurang)
- Siap deploy ke Streamlit Cloud
"""

import streamlit as st
from docx import Document
from io import BytesIO
from datetime import datetime
import os
import re

# ====================== PENGATURAN ======================
PASSWORD = "Bnpb123#"          
TEMPLATE_PATH = "template_surat.docx"
# ========================================================


def check_password():
    """Simple password gate"""
    def password_entered():
        if st.session_state.get("password") == PASSWORD:
            st.session_state["authenticated"] = True
            if "password" in st.session_state:
                del st.session_state["password"]
        else:
            st.session_state["authenticated"] = False

    if "authenticated" not in st.session_state:
        st.session_state["authenticated"] = False

    if not st.session_state["authenticated"]:
        st.title("🔒 Generator IL Darurat Wilayah I BNPB")
        st.caption("Developed by Rifqi Adlian Syah")
        st.text_input(
            "Masukkan Password",
            type="password",
            on_change=password_entered,
            key="password"
        )
        if st.session_state.get("authenticated") is False and "password" in st.session_state:
            st.error("Password salah")
        st.stop()


def fill_template(context: dict) -> bytes:
    """Isi template Word dengan data dari form"""
    if not os.path.exists(TEMPLATE_PATH):
        st.error(f"Template tidak ditemukan: {TEMPLATE_PATH}")
        st.stop()

    doc = Document(TEMPLATE_PATH)

    def replace_in_paragraph(p, context):
        text = p.text
        new_text = text
        for key, val in context.items():
            placeholder = "{{ " + key + " }}"
            if placeholder in new_text:
                new_text = new_text.replace(placeholder, str(val))
        if new_text != text:
            if p.runs:
                p.runs[0].text = new_text
                for r in p.runs[1:]:
                    r.text = ""
            else:
                p.add_run(new_text)

    for p in doc.paragraphs:
        replace_in_paragraph(p, context)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    replace_in_paragraph(p, context)

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def init_dasar_list():
    """Inisialisasi list dasar hukum di session_state"""
    if "dasar_list" not in st.session_state:
        st.session_state.dasar_list = [
            "Instruksi Presiden Nomor 3 Tahun 2020, tanggal 28 Februari 2020 tentang Penanggulangan Kebakaran Hutan dan Lahan;",
            "Keputusan Gubernur .......... Nomor .......... tanggal .......... tentang Penetapan Status Siaga Darurat ...;",
            "Surat Gubernur .......... Nomor .......... tanggal .......... perihal Permohonan Bantuan Helikopter Waterbombing;",
        ]


# ====================== UI ======================
st.set_page_config(
    page_title="Generator Surat BNPB",
    page_icon="📄",
    layout="wide"
)

check_password()
init_dasar_list()

st.title("📄 Generator Surat Penunjukan / IL Perpanjangan Darurat Wilayah 1 (BNPB)")
st.caption("Isi form → Generate → Download Word")

# ----- Form Informasi Surat & Penerima -----
col1, col2 = st.columns(2)

with col1:
    st.subheader("1. Informasi Surat")
    nomor_surat = st.text_input("Nomor Surat", value="B- XXX/KA BNPB/PD.01.04/09/2026")

    bulan_map = {
        1: "Januari", 2: "Februari", 3: "Maret", 4: "April",
        5: "Mei", 6: "Juni", 7: "Juli", 8: "Agustus",
        9: "September", 10: "Oktober", 11: "November", 12: "Desember"
    }
    now = datetime.now()
    default_tanggal = f"{now.day} {bulan_map[now.month]} {now.year}"
    tanggal = st.text_input("Tanggal", value=default_tanggal)

    perihal = st.text_area(
        "Perihal (Hal)",
        value="Penunjukan Pelaksana Kegiatan Pengadaan Jasa Lainnya berupa Jasa Angkutan Udara (Sewa Helikopter Water Bombing) dalam rangka Penanganan Siaga Darurat akibat Kebakaran Hutan dan Lahan di Wilayah Provinsi ... Tahun 2026",
        height=110
    )

with col2:
    st.subheader("2. Penerima")

    col_a, col_b = st.columns([1.4, 3])
    with col_a:
        st.markdown(
            "<p style='margin-top: 0.5rem;'><b>Yth.</b> Direktur PT.</p>",
            unsafe_allow_html=True
        )
    with col_b:
        nama_perusahaan = st.text_input(
            "Nama Perusahaan",
            value="",
            placeholder="Nama Perusahaan",
            label_visibility="collapsed"
        )

    yth_penerima = f"Direktur PT. {nama_perusahaan}".strip()
    nama_PT = f"PT. {nama_perusahaan}".strip().title() if nama_perusahaan else "PT."

    alamat = st.text_area(
        "Alamat Lengkap",
        value="Jl. ..................\nKota .................., Kode Pos",
        height=90
    )
    

# ----- Dasar Hukum DINAMIS -----
st.subheader("3. Dasar Hukum / Referensi")
st.caption("Bisa ditambah atau dikurangi sesuai kebutuhan. Kosongkan yang tidak dipakai.")

# Tombol tambah
col_btn1, col_btn2, _ = st.columns([1, 1, 4])
with col_btn1:
    if st.button("➕ Tambah Dasar", use_container_width=True):
        st.session_state.dasar_list.append("")
        st.rerun()

with col_btn2:
    if st.button("🗑️ Hapus Semua", use_container_width=True):
        st.session_state.dasar_list = [""]
        st.rerun()

# Tampilkan input untuk setiap dasar
new_list = []
for i, item in enumerate(st.session_state.dasar_list):
    c1, c2 = st.columns([10, 1])
    with c1:
        val = st.text_input(
            f"Dasar {i+1}",
            value=item,
            key=f"dasar_input_{i}",
            label_visibility="collapsed" if i > 0 else "visible"
        )
        # Show label only for first, others use placeholder style
        if i == 0:
            pass  # label already shown
        new_list.append(val)
    with c2:
        if len(st.session_state.dasar_list) > 1:
            if st.button("✖", key=f"hapus_{i}", help="Hapus baris ini"):
                st.session_state.dasar_list.pop(i)
                st.rerun()

st.session_state.dasar_list = new_list

# ----- Detail Kegiatan -----
st.subheader("4. Detail Kegiatan & Penyedia")
c3, c4, c5 = st.columns(3)
with c3:
    # nama_penyedia = st.text_input("Nama Penyedia (PT)", value="PT. ..................")
    nama_penyedia = nama_PT
    jumlah_unit = st.text_input("Jumlah Unit", value="1 (satu)")
with c4:
    jenis_kegiatan = st.text_input("Jenis Kegiatan", value="helikopter untuk melaksanakan kegiatan water bombing")
    spek_helikopter = st.text_input("Spesifikasi Helikopter", value="helikopter water bombing Tipe Sikorsky S-61A, Reg. NXXXX")
with c5:
    wilayah = st.selectbox(
        "Wilayah Operasi",
        options=[
            "Provinsi Riau",
            "Provinsi Kalimantan Tengah",
            "Provinsi Aceh",
            "Provinsi Kepulauan Riau",
            # tambah lagi di sini
        ],
        index=0
    )
    tahun_sekarang = datetime.now().year

    tahun = st.selectbox(
        "Tahun",
        options=[f"Tahun {y}" for y in range(tahun_sekarang-5, tahun_sekarang+5)],
        index=5
    )

st.divider()

if st.button("🚀 Generate Surat Word", type="primary", use_container_width=True):
    # Bangun daftar dasar sebagai teks bernomor
    dasar_items = [d.strip() for d in st.session_state.dasar_list if d.strip()]
    if not dasar_items:
        st.warning("Minimal isi 1 Dasar Hukum.")
        st.stop()

    daftar_dasar_lines = []
    for idx, text in enumerate(dasar_items, start=1):
        # Pastikan diakhiri titik koma kalau belum
        if not text.endswith((";", ".")):
            text = text + ";"
        daftar_dasar_lines.append(f"{idx}. {text}")

    daftar_dasar = "\n".join(daftar_dasar_lines)

    context = {
        "nomor_surat": nomor_surat,
        "tanggal": tanggal,
        "perihal": perihal,
        "yth_penerima": yth_penerima,
        "alamat": alamat,
        "daftar_dasar": daftar_dasar,
        "nama_penyedia": nama_penyedia,
        "jumlah_unit": jumlah_unit,
        "jenis_kegiatan": jenis_kegiatan,
        "spek_helikopter": spek_helikopter,
        "wilayah": wilayah,
        "tahun": tahun,
    }

    with st.spinner("Sedang generate surat..."):
        try:
            file_bytes = fill_template(context)
            st.success(f"✅ Surat berhasil dibuat! ({len(dasar_items)} dasar hukum)")

            safe_name = nomor_surat.replace("/", "-").replace(" ", "").replace(":", "")
            st.download_button(
                label="📥 Download Surat Word (.docx)",
                data=file_bytes,
                file_name=f"Surat_Penunjukan_{safe_name}.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                use_container_width=True
            )
        except Exception as e:
            st.error(f"Gagal generate: {e}")

st.markdown("---")
st.info("💡 **Tips:** Setelah download, buka file di Microsoft Word. Tolong dicek lagi kerapihan file sebelum dikirim karena isi file dinamis")
st.info("Developed by Rifqi Adlian Syah")